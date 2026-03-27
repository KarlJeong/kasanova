import io

import fitz
import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.rag.parsers import parse_by_extension, UnsupportedFormatError
from app.rag.parsers.pdf_parser import parse_pdf
from app.rag.parsers.pptx_parser import parse_pptx
from app.rag.parsers.docx_parser import parse_docx


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """1페이지짜리 PDF를 인메모리로 생성한다."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF World")
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture
def sample_pptx_bytes() -> bytes:
    """1슬라이드짜리 PPTX를 인메모리로 생성한다."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Hello PPTX World"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """1단락짜리 DOCX를 인메모리로 생성한다."""
    doc = DocxDocument()
    doc.add_paragraph("Hello DOCX World")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestPdfParser:
    def test_extract_text(self, sample_pdf_bytes: bytes) -> None:
        text = parse_pdf(sample_pdf_bytes)
        assert "Hello PDF World" in text

    def test_empty_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_pdf(b"")

    def test_corrupt_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_pdf(b"not a pdf file")


class TestPptxParser:
    def test_extract_text(self, sample_pptx_bytes: bytes) -> None:
        text = parse_pptx(sample_pptx_bytes)
        assert "Hello PPTX World" in text

    def test_empty_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_pptx(b"")

    def test_corrupt_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_pptx(b"not a pptx file")


class TestDocxParser:
    def test_extract_text(self, sample_docx_bytes: bytes) -> None:
        text = parse_docx(sample_docx_bytes)
        assert "Hello DOCX World" in text

    def test_empty_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_docx(b"")

    def test_corrupt_bytes_raises(self) -> None:
        with pytest.raises(ValueError):
            parse_docx(b"not a docx file")


class TestParseByExtension:
    def test_dispatch_pdf(self, sample_pdf_bytes: bytes) -> None:
        text = parse_by_extension(".pdf", sample_pdf_bytes)
        assert "Hello PDF World" in text

    def test_dispatch_pptx(self, sample_pptx_bytes: bytes) -> None:
        text = parse_by_extension(".pptx", sample_pptx_bytes)
        assert "Hello PPTX World" in text

    def test_dispatch_docx(self, sample_docx_bytes: bytes) -> None:
        text = parse_by_extension(".docx", sample_docx_bytes)
        assert "Hello DOCX World" in text

    def test_unsupported_extension_raises(self) -> None:
        with pytest.raises(UnsupportedFormatError):
            parse_by_extension(".txt", b"some text")
