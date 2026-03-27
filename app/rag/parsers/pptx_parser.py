import io

from pptx import Presentation


def parse_pptx(data: bytes) -> str:
    """PPTX 바이트에서 슬라이드별 텍스트를 추출한다."""
    if not data:
        raise ValueError("빈 PPTX 데이터입니다.")
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ValueError(f"PPTX 파싱 실패: {e}") from e

    slides: list[str] = []
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            slides.append("\n".join(texts))

    if not slides:
        raise ValueError("PPTX에서 텍스트를 추출할 수 없습니다.")
    return "\n\n".join(slides)
